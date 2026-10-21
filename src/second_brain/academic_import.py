"""Importación revisable de listas y horarios; fuentes privadas para FTS/RAG."""
import csv,hashlib,io,json,os,re,subprocess,sys,zipfile
from pathlib import Path
from .core import Brain,identifier,now,normal
FIELDS={'lista':['matricula','nombre'],'horario':['materia','grupo','aula','inicio','fin']}
ALIASES={'numero':'matricula','matricula':'matricula','nombre completo':'nombre','alumno':'nombre','clave':'materia','asignatura':'materia','salon':'aula','fecha inicio':'inicio','fecha fin':'fin'}
def normalize_rows(rows,kind):
 if len(rows)>1001:raise ValueError('Máximo 1000 filas por archivo')
 if not rows:return []
 headers=[ALIASES.get(normal(str(v)).strip(),normal(str(v)).strip()) for v in rows[0]]
 if len(set(headers))!=len(headers):raise ValueError('Encabezados duplicados')
 result=[]
 for n,row in enumerate(rows[1:],2):
  if not any(str(v).strip() for v in row):continue
  item={h:str(row[i] if i<len(row) and row[i] is not None else '').strip() for i,h in enumerate(headers)}
  result.append({k:item.get(k,'') for k in FIELDS[kind]})
 return result
def preview(path,extension,kind,sheet_name=None):
 if kind not in FIELDS:raise ValueError('Tipo de importación inválido')
 path=Path(path)
 if path.stat().st_size>10*1024**2:raise ValueError('Máximo 10 MB')
 text='';warnings=[];sheets=[];selected_sheet='';metadata={}
 if extension=='csv':
  data=path.read_bytes()
  try:text=data.decode('utf-8-sig')
  except UnicodeDecodeError:text=data.decode('cp1252')
  try:delimiter=csv.Sniffer().sniff(text[:8192],delimiters=',;\t').delimiter
  except csv.Error:delimiter=';'
  rows=list(csv.reader(io.StringIO(text),delimiter=delimiter))
 elif extension in {'xls','xlsx'}:
  if extension=='xls':
   import xlrd
   book=xlrd.open_workbook(str(path));sheets=book.sheet_names();selected_sheet=sheet_name or sheets[0];sheet=book.sheet_by_name(selected_sheet)
   if sheet.nrows>1001 or sheet.ncols>40:raise ValueError('Hoja demasiado grande')
   rows=[[sheet.cell_value(r,c) for c in range(sheet.ncols)] for r in range(sheet.nrows)]
   rows=[[str(int(v)) if isinstance(v,float) and v.is_integer() else v for v in row] for row in rows]
  else:
   with zipfile.ZipFile(path) as archive:
    if sum(i.file_size for i in archive.infolist())>40*1024**2:raise ValueError('Excel expandido demasiado grande')
   import openpyxl
   book=openpyxl.load_workbook(io.BytesIO(path.read_bytes()),read_only=True,data_only=True)
   sheets=book.sheetnames;selected_sheet=sheet_name or sheets[0]
   if selected_sheet not in sheets:book.close();raise ValueError('Hoja inexistente')
   sheet=book[selected_sheet]
   if (sheet.max_row or 0)>1001 or (sheet.max_column or 0)>40:book.close();raise ValueError('Hoja demasiado grande')
   rows=[]
   for row in sheet.values:
    if len(rows)>=1001 or len(row)>40:book.close();raise ValueError('Hoja demasiado grande')
    rows.append([str(int(v)) if isinstance(v,(int,float)) and not isinstance(v,bool) and float(v).is_integer() else v for v in row])
   book.close()
  warnings.append('Selecciona la hoja que corresponde al grupo. Matrículas con ceros iniciales deben estar guardadas como texto.')
 elif extension=='pdf' and kind=='horario':
  text=subprocess.run(['pdftotext','-layout',str(path),'-'],capture_output=True,text=True,check=True,timeout=30).stdout
  if not text.strip():raise ValueError('PDF sin texto: requiere OCR o convertir a CSV')
  rows=[re.split(r'\s{2,}',line.strip()) for line in text.splitlines() if line.strip()]
  warnings.append('PDF: revisa las columnas extraídas. Si no tiene una tabla reconocible, transcribe las filas en la vista previa; no se infieren fechas ni grupos.')
 else:raise ValueError('Formato no permitido')
 if kind=='lista':
  header=None
  for i,row in enumerate(rows[:100]):
   keys=[normal(str(v or '')).strip().rstrip(':') for v in row]
   if 'matricula' in keys and ('nombre' in keys or 'nombre completo' in keys):header=i;break
   if keys and keys[0] in {'materia','grupo','maestro','docente'}:
    metadata[keys[0]]=next((str(v) for v in row[1:] if v is not None and str(v).strip()),'')
  if header is None:raise ValueError('No se encontraron columnas Matrícula y Nombre en esta hoja')
  rows=rows[header:]
  if metadata.get('grupo') and selected_sheet and normal(metadata['grupo']).replace('-','').replace('_','')!=normal(selected_sheet).replace('-','').replace('_',''):
   warnings.append('El grupo dentro de la hoja ('+metadata['grupo']+') difiere del nombre de la hoja ('+selected_sheet+'). Confirma el grupo correcto antes de registrar.')
 parsed=normalize_rows(rows,kind)
 if kind=='lista':
  parsed=[r for r in parsed if r['matricula'] or r['nombre']]
  if not parsed:raise ValueError('La hoja no contiene alumnos')
 if extension=='pdf' and not all(k in [normal(str(x)) for x in rows[0]] for k in FIELDS[kind]):parsed=[]
 if any(not all(r.values()) for r in parsed):warnings.append('Hay campos vacíos: complétalos antes de registrar.')
 return {'rows':parsed,'text':text[:60000],'warnings':warnings,'fields':FIELDS[kind],'sheets':sheets,'sheet':selected_sheet,'metadata':metadata}
def archive(brain,path,extension,kind,name,records,context):
 if kind not in FIELDS or not records or len(records)>1000:raise ValueError('Registros inválidos')
 source=Path(path).read_bytes();digest=hashlib.sha256(source+json.dumps([kind,records,context],sort_keys=True).encode()).hexdigest()
 found=brain.rows('SELECT id FROM documents WHERE hash=?',(digest,))
 if found:return {'document_id':found[0]['id'],'duplicate':True}
 docid=identifier();folder=brain.home/'data/sources';folder.mkdir(exist_ok=True)
 original=folder/(digest+'.'+extension);original.write_bytes(source);os.chmod(original,0o600)
 payload={'tipo':kind,'fuente':name,'fecha':now(),'contexto':context,'registros':records,'original':original.name,'estado':context.get('estado_lista','registrado')}
 target=folder/(digest+'.json');target.write_text(json.dumps(payload,ensure_ascii=False,indent=2));os.chmod(target,0o600)
 with brain.db:
  brain.db.execute('INSERT INTO documents VALUES(?,?,?,?,?,?,?)',(docid,digest,name+' · '+kind,'ACADEMICO',str(target),1,now()))
  for row in records:
   text=json.dumps({'tipo':kind,'contexto':context,**row},ensure_ascii=False);cid=identifier()
   brain.db.execute('INSERT INTO chunks VALUES(?,?,?,?,?)',(cid,docid,1,text,hashlib.sha256(text.encode()).hexdigest()))
   brain.db.execute('INSERT INTO search VALUES(?,?,?)',(cid,'document',text))
 return {'document_id':docid,'rows':len(records),'duplicate':False}
if __name__=='__main__':
 try:
  args=json.load(sys.stdin)
  if args['action']=='current_weekly':
   from datetime import datetime,time
   from zoneinfo import ZoneInfo
   b=Brain(args['home']);result=[]
   try:
    for d in b.rows("SELECT path FROM documents WHERE career='ACADEMICO'"):
     content=json.loads(Path(d['path']).read_text());ctx=content.get('contexto',{})
     if ctx.get('docente_id')!=args['teacherId'] or ctx.get('horario_tipo')!='semanal':continue
     current=datetime.fromtimestamp(args['timestamp'],ZoneInfo(ctx.get('zona_horaria','America/Mexico_City')))
     if ctx.get('periodo')=='AGO-DIC 2026' and (current.year!=2026 or current.month not in range(8,13)):continue
     for block in content['registros']:
      if block['dia_semana']!=current.weekday():continue
      start=datetime.combine(current.date(),time.fromisoformat(block['hora_inicio']),current.tzinfo)
      end=datetime.combine(current.date(),time.fromisoformat(block['hora_fin']),current.tzinfo)
      if start<=current<end:result.append({**block,'starts':int(start.timestamp()),'ends':int(end.timestamp()),'periodo':ctx['periodo']})
   finally:b.close()
  elif args['action']=='monthly_weekly':
   from datetime import date,datetime,time
   from calendar import monthrange
   from zoneinfo import ZoneInfo
   year,month=map(int,args['month'].split('-'));result=[];b=Brain(args['home'])
   try:
    for d in b.rows("SELECT path FROM documents WHERE career='ACADEMICO'"):
     content=json.loads(Path(d['path']).read_text());ctx=content.get('contexto',{})
     if ctx.get('docente_id')!=args['teacherId'] or ctx.get('horario_tipo')!='semanal':continue
     if ctx.get('periodo')=='AGO-DIC 2026' and (year!=2026 or month not in range(8,13)):continue
     for n in range(1,monthrange(year,month)[1]+1):
      day=date(year,month,n)
      for block in content['registros']:
       if block['dia_semana']!=day.weekday():continue
       tz=ZoneInfo(ctx.get('zona_horaria','America/Mexico_City'))
       begin=datetime.combine(day,time.fromisoformat(block['hora_inicio']),tz)
       end=datetime.combine(day,time.fromisoformat(block['hora_fin']),tz)
       result.append({**block,'starts':int(begin.timestamp()),'ends':int(end.timestamp()),'periodo':ctx['periodo'],'date':day.isoformat()})
   finally:b.close()
  elif args['action']=='expand_weekly':
   from datetime import date,datetime,timedelta,time
   from zoneinfo import ZoneInfo
   b=Brain(args['home'])
   try:
    documents=b.rows('SELECT path,name FROM documents WHERE id=?',(args['documentId'],))
    if not documents:raise ValueError('Fuente de horario inexistente')
    path=Path(documents[0]['path'])
    if not path.resolve().is_relative_to((b.home/'data/sources').resolve()):raise ValueError('Fuente inválida')
    content=json.loads(path.read_text());ctx=content.get('contexto',{})
    if ctx.get('docente_id')!=args['teacherId'] or ctx.get('horario_tipo')!='semanal':raise ValueError('Horario de otro docente')
    start=date.fromisoformat(args['start']);end=date.fromisoformat(args['end'])
    if end<start or (end-start).days>180:raise ValueError('Rango de fechas inválido (máximo 180 días)')
    if ctx.get('periodo')=='AGO-DIC 2026' and (start<date(2026,8,1) or end>date(2026,12,31)):raise ValueError('El horario corresponde a AGO-DIC 2026')
    exclude={date.fromisoformat(d) for d in args.get('exclude',[])};rows=[];day=start
    while day<=end:
     if day not in exclude:
      for block in content['registros']:
       if day.weekday()!=block['dia_semana']:continue
       begin=datetime.combine(day,time.fromisoformat(block['hora_inicio']),ZoneInfo(ctx['zona_horaria']))
       finish=datetime.combine(day,time.fromisoformat(block['hora_fin']),ZoneInfo(ctx['zona_horaria']))
       rows.append({'materia':block['materia'],'grupo':block['grupo'],'aula':args['room'],'inicio':begin.isoformat(),'fin':finish.isoformat()})
     day+=timedelta(days=1)
    if not rows:raise ValueError('No hay clases en ese rango')
    if len(rows)>1000:raise ValueError('Demasiadas sesiones')
    original=b.home/'data/sources'/content['original']
    if not original.resolve().is_relative_to((b.home/'data/sources').resolve()):raise ValueError('Original inválido')
    result={'rows':rows,'path':str(original),'name':documents[0]['name']+' · '+str(start)+'–'+str(end)}
   finally:b.close()
  elif args['action']=='teacher_sources':
   b=Brain(args['home']);result=[]
   try:
    for d in b.rows("SELECT id,path,name FROM documents WHERE career='ACADEMICO'"):
     content=json.loads(Path(d['path']).read_text());ctx=content.get('contexto',{})
     if ctx.get('docente_id')==args['teacherId']:
      result.append({'id':d['id'],'name':d['name'],'context':ctx,'kind':content.get('tipo'),'rows':len(content.get('registros',[])),'weekly':content.get('registros',[]) if ctx.get('horario_tipo')=='semanal' else [],'roster':content.get('registros',[]) if content.get('tipo')=='lista' else []})
   finally:b.close()
  elif args['action']=='preview':result=preview(args['path'],args['extension'],args['kind'],args.get('sheet'))
  else:
   b=Brain(args['home'])
   try:result=archive(b,args['path'],args['extension'],args['kind'],args['name'],args['records'],args['context'])
   finally:b.close()
  print(json.dumps(result,ensure_ascii=False))
 except Exception as exc:
  import traceback;traceback.print_exc(file=sys.stderr)
  print(json.dumps({'error':str(exc)},ensure_ascii=False));sys.exit(1)
